# -*- coding: utf-8 -*-
"""거래내역 1/2 세로변환 GUI · Python 3.9+ / openpyxl"""
import json
import queue
import re
import threading
import traceback
import zipfile
import os
import sys
import tempfile
import copy
import warnings
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, colorchooser
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    WindowBase = TkinterDnD.Tk
    DRAG_DROP_AVAILABLE = True
except ImportError:
    WindowBase = tk.Tk
    DRAG_DROP_AVAILABLE = False
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

APP = '전송형 전자영장 집행포털 [입출금내역제공내역 일괄변환 프로그램]'
APP_DIR = Path(sys.executable if getattr(sys, 'frozen', False) else __file__).resolve().parent
CONFIG = APP_DIR / 'TransactionConverter_Settings.json'
OLD_CONFIG = Path.home() / '.transaction_vertical_converter_v2.json'
# 출력 논리열: 1~53 (A:BA) + 86~90 (CH:CL) = 58개
LOGICAL_COLS = list(range(1, 54)) + list(range(86, 91))
LOAN_LABELS = {43:'대출금지급 거래금액',44:'대출원금상환 거래금액',45:'대출원금 적요',46:'대출 상대방 계좌주명',47:'대출 상대방금융기관코드',48:'대출 상대방 계좌번호',49:'대출이자 거래금액',50:'대출이자 적요',51:'대출이자 상대방 계좌주명',52:'대출이자 상대방금융기관코드',53:'대출이자 상대방 계좌번호'}
DEFAULT_COLS = [11, 12] + list(range(22, 40))  # K,L,V:AM
DEFAULT_BG = {'입금': '#D4F0E6', '지급': '#FFE0B8'}
DEFAULT_FG = {'입금': '#075448', '지급': '#8F3400'}
NS = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

# 색상 조합 프리셋 (적용하지 않는 거래는 검정 글자/흰 배경)
COLOR_PRESETS = {
    '기본 · 지급만 연주황': {'입금': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}, '지급': {'enabled': True, 'bg': '#FFE0B8', 'fg': '#8F3400'}},
    '지급만 연분홍': {'입금': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}, '지급': {'enabled': True, 'bg': '#FCE1E4', 'fg': '#9F1D35'}},
    '지급만 연회색': {'입금': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}, '지급': {'enabled': True, 'bg': '#DFE2E6', 'fg': '#23272B'}},
    '입금만 연녹색': {'입금': {'enabled': True, 'bg': '#D4F0E6', 'fg': '#075448'}, '지급': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}},
    '입금만 연파랑': {'입금': {'enabled': True, 'bg': '#DCEBFF', 'fg': '#154B83'}, '지급': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}},
    '입금 녹색 / 지급 주황': {'입금': {'enabled': True, 'bg': '#D4F0E6', 'fg': '#075448'}, '지급': {'enabled': True, 'bg': '#FFE0B8', 'fg': '#8F3400'}},
    '입금 파랑 / 지급 분홍': {'입금': {'enabled': True, 'bg': '#DCEBFF', 'fg': '#154B83'}, '지급': {'enabled': True, 'bg': '#FCE1E4', 'fg': '#9F1D35'}},
    '입금 회색 / 지급 노랑': {'입금': {'enabled': True, 'bg': '#E4E7EB', 'fg': '#252B34'}, '지급': {'enabled': True, 'bg': '#FFF1B6', 'fg': '#805000'}},
    '모두 색상 없음': {'입금': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}, '지급': {'enabled': False, 'bg': '#FFFFFF', 'fg': '#000000'}},
}
DEFAULT_COLOR_PRESET = '기본 · 지급만 연주황'

# 가운데 정렬 대상: 거래 날짜/시간·코드·금액 및 지정 계좌/거래 정보.
# 원본 머리글의 공백 및 개행과 무관하게 항목명으로 판정한다.
CENTER_ALIGN_HEADERS = {
    '입출금일시', '거래일자', '시간', '거래코드', '전체거래금액',
    '상대방금융기관코드',
    '조회시작', '조회종료', '조회시작시각', '조회종료시각', '종목코드', '계좌번호',
    '계좌번호(신)', '계좌번호(구)',
    '평생계좌번호', '거래매체명', '거래점코드', '자금내역',
}

def normalized_header(value):
    return re.sub(r'\s+', '', str(value or ''))

def next_output_path(path):
    """동일 폴더에 원본_변환.xlsx, 원본_변환2.xlsx, ... 순서로 생성."""
    path = Path(path)
    candidate = path.with_name(path.stem + '_변환.xlsx')
    i = 2
    while candidate.exists():
        candidate = path.with_name(f'{path.stem}_변환{i}.xlsx')
        i += 1
    return candidate


def colname(n):
    out = ''
    while n:
        n, r = divmod(n-1, 26)
        out = chr(65+r)+out
    return out

def get_xml_sheet_header(path):
    """읽기 전용 ZIP/XML 파싱: 서식/셀 변경 없음. 3행 명칭, 최대 행 수 검색."""
    labels = {i: colname(i) for i in range(1,91)}
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read('xl/workbook.xml'))
        rns = {'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        rid = wb.find('x:sheets/x:sheet', NS).get('{'+rns['r']+'}id')
        relationships = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target = next(rel.get('Target') for rel in relationships if rel.get('Id') == rid)
        if target.startswith('/'):
            member = target.lstrip('/')
        else:
            import posixpath
            member = posixpath.normpath(posixpath.join('xl',target))
        shared = []
        if 'xl/sharedStrings.xml' in z.namelist():
            strings = ET.fromstring(z.read('xl/sharedStrings.xml'))
            shared = [''.join(t.text or '' for t in si.iter('{'+NS['x']+'}t')) for si in strings]
        # Stream parse large files without storing all cells
        maxrow = 0
        with z.open(member) as fp:
            for event, elem in ET.iterparse(fp, events=('end',)):
                if elem.tag == '{'+NS['x']+'}row':
                    rn = int(elem.get('r', '0'))
                    maxrow = max(maxrow, rn)
                    if rn == 3:
                        for c in elem.findall('x:c', NS):
                            match = re.match(r'([A-Z]+)', c.get('r', ''))
                            if not match: continue
                            idx = 0
                            for ch in match.group(1): idx = idx*26+ord(ch)-64
                            if idx not in labels: continue
                            if c.get('t') == 'inlineStr':
                                name = ''.join(t.text or '' for t in c.iter('{'+NS['x']+'}t'))
                            else:
                                v = c.find('x:v',NS)
                                name = (shared[int(v.text)] if c.get('t')=='s' and v is not None else (v.text or '' if v is not None else ''))
                            if name.strip(): labels[idx] = name.strip()
                    elem.clear()
    return labels, maxrow

def load_settings():
    opts = {'colors': copy.deepcopy(COLOR_PRESETS[DEFAULT_COLOR_PRESET]),
            'color_preset': DEFAULT_COLOR_PRESET,
            'presets': {'기본 프리셋1 (거래내역만-대출제외)': DEFAULT_COLS},
            'selected': '기본 프리셋1 (거래내역만-대출제외)', 'skip_blank': True}
    try:
        path_to_read = CONFIG if CONFIG.exists() else OLD_CONFIG
        user = json.loads(path_to_read.read_text(encoding='utf-8'))
        if isinstance(user, dict):
            opts.update(user)
            opts['presets']['기본 프리셋1 (거래내역만-대출제외)'] = DEFAULT_COLS
            # v3.2 이하 설정은 enabled가 없으므로 새 기본값인 지급만 색상 적용
            for kind in ('입금', '지급'):
                opts['colors'].setdefault(kind, {})
                opts['colors'][kind].setdefault('enabled', kind == '지급')
            opts.setdefault('color_preset', '사용자 지정')
    except (OSError, ValueError, TypeError):
        pass
    if opts.get('selected') not in opts['presets']:
        opts['selected']='기본 프리셋1 (거래내역만-대출제외)'
    return opts

def save_settings(opts):
    CONFIG.write_text(json.dumps(opts,ensure_ascii=False,indent=2),encoding='utf-8')

def output_order(selected):
    """원본 입출금일시 바로 뒤 거래일자·시간을 생성하고 거래코드를 시간 뒤로 이동."""
    cols=[n for n in LOGICAL_COLS if n in selected and n!=22]
    if 31 in cols:
        ix=cols.index(31)+1
        cols[ix:ix]=[91,92]  # 가상 열 (거래일자, 시간)
    if 22 in selected:
        anchor=92 if 92 in cols else (31 if 31 in cols else None)
        pos=cols.index(anchor)+1 if anchor is not None else 0
        cols.insert(pos,22)
    return cols

def derive_datetime(value, part):
    digits=re.sub(r'[^0-9]','',str(value or ''))
    if len(digits)<12:return ''
    if part==91:return digits[:8]
    return digits[8:10]+':'+digits[10:12]+(':'+digits[12:14] if len(digits)>=14 else ':00')

def date_only(value):
    """조회시각에서 날짜만 반환. 14자리/8자리/날짜 객체를 지원."""
    if value is None or str(value).strip() == '':
        return ''
    if hasattr(value, 'strftime'):
        try:
            return value.strftime('%Y-%m-%d')
        except (ValueError, TypeError):
            pass
    raw = str(value).strip()
    digits = re.sub(r'[^0-9]', '', raw)
    if len(digits) >= 8:
        return f'{digits[:4]}-{digits[4:6]}-{digits[6:8]}'
    return raw

def usable_tx_code(val):
    code=str(val or '').strip()
    return code not in ('', '0', '0.0')

def process_file(path, selected, colors, skip_blank, report):
    """openpyxl만으로 XLSX 생성. OOXML 수동 수정 없음."""
    path = Path(path)
    dest = next_output_path(path)
    # 원본 금융기관 XLSX에 기본 스타일 정의가 없는 경우에만 발생하는 안내 경고를 제한적으로 제외.
    # 원본 내용을 수정하거나 다른 경고를 숨기지는 않음.
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message=r'Workbook contains no default style, apply openpyxl\'s default', category=UserWarning, module=r'openpyxl\.styles\.stylesheet')
        wb = load_workbook(path)
    src = wb.worksheets[0]
    if '변환' in wb.sheetnames or ('원본시트' in wb.sheetnames and src.title != '원본시트'):
        raise ValueError('이미 변환된 파일입니다. 원본 파일을 선택하세요.')
    source_labels, _ = get_xml_sheet_header(path)
    # 원본시트는 원본 셀의 값과 머리글을 그대로 보존한다.
    src.title = '원본시트'
    src.freeze_panes = 'A4'
    src.auto_filter.ref = f'A3:CL{max(3,src.max_row)}'
    out = wb.create_sheet('변환(요약)',0)
    full = wb.create_sheet('변환(원본칼럼 유지)',1)
    ordered = output_order(selected)
    if not ordered: raise ValueError('출력할 열을 하나 이상 선택하세요.')

    def two_line(t):
        t=str(t or '').strip()
        if len(t)<6:return t
        if ' ' in t:
            parts=t.split(' '); mid=max(1,len(parts)//2)
            return ' '.join(parts[:mid])+'\n'+' '.join(parts[mid:])
        mid=(len(t)+1)//2
        return t[:mid]+'\n'+t[mid:]

    labels=[('거래일자' if n==91 else '시간' if n==92 else '조회시작' if n==11 else '조회종료' if n==12 else LOAN_LABELS.get(n,source_labels.get(n,colname(n)))) for n in ordered]
    for i,label in enumerate(labels,1):
        out.cell(1,i,two_line(label))
    out.row_dimensions[1].height=30
    out.freeze_panes='A2'
    # 전체 시트: 원본 A~CL 90개 열, 상단 1~3행과 병합 구조를 보존한다.
    for original_row in src.iter_rows(min_row=1,max_row=3,max_col=90):
        for original_cell in original_row:
            dest_cell=full.cell(original_cell.row,original_cell.column)
            if isinstance(original_cell,MergedCell): continue
            dest_cell.value=original_cell.value
            if original_cell.has_style:
                dest_cell._style=copy.copy(original_cell._style)
            if original_cell.number_format:dest_cell.number_format=original_cell.number_format
            dest_cell.alignment=copy.copy(original_cell.alignment)
    for merge in src.merged_cells.ranges:
        if merge.max_row<=3 and merge.max_col<=90:
            full.merge_cells(str(merge))
    for r in range(1,4):
        if src.row_dimensions[r].height is not None:full.row_dimensions[r].height=src.row_dimensions[r].height
    for ci in range(1,91):
        letter=get_column_letter(ci)
        source_dim=src.column_dimensions.get(letter)
        if source_dim is not None and source_dim.width is not None:
            full.column_dimensions[letter].width=source_dim.width
    full.freeze_panes='A4'
    full.auto_filter.ref='A3:CL3'
    full.sheet_view.showGridLines=src.sheet_view.showGridLines
    full_count=0
    def append_full(source_row,offset):
        nonlocal full_count
        full_count+=1
        target_row=full_count+3
        for ci in range(1,91):
            original_ci = ci + offset if 22<=ci<=53 else ci
            is_empty_right = 54<=ci<=85
            original_cell=src.cell(source_row,original_ci)
            target=full.cell(target_row,ci)
            target.value=None if is_empty_right else original_cell.value
            if original_cell.has_style: target._style=copy.copy(original_cell._style)
            target.font=copy.copy(original_cell.font)
            updated_font=copy.copy(target.font)
            updated_font.bold=False
            target.font=updated_font
            target.alignment=copy.copy(original_cell.alignment)
        dim=src.row_dimensions[source_row]
        if dim.height is not None:full.row_dimensions[target_row].height=dim.height

    count=0
    for rn in range(4,src.max_row+1):
        vals=[src.cell(rn,c).value for c in range(1,91)]
        for offset in (0,32):
            if offset==32 and not usable_tx_code(vals[53]):
                continue
            tx=vals[21+offset:53+offset]
            if skip_blank and not any(v is not None and str(v).strip() for v in tx):
                continue
            append_full(rn,offset)
            row=[]
            for logical in ordered:
                v=(derive_datetime(vals[30+offset],logical) if logical in (91,92)
                   else vals[(logical+offset if 22<=logical<=53 else logical)-1])
                if logical in (26,27) and v is not None:
                    raw=str(v).strip()
                    if raw and set(raw)=={'0'}:v=' '
                if logical in (11,12):
                    v=date_only(v)
                if re.sub(r'\s+','',labels[len(row)]) in ('전체거래금액','현금거래금액','잔액'):
                    if isinstance(v,str):
                        raw=v.replace(',','').strip()
                        if re.fullmatch(r'[+-]?\d+(?:\.\d+)?',raw):
                            v=float(raw) if '.' in raw else int(raw)
                row.append(v)
            out.append(row)
            count+=1
        if rn%250==0 or rn==src.max_row:
            report(f'{path.name}: 원본 {rn:,}/{src.max_row:,}행 · 거래 {count:,}건')

    width=len(ordered)
    money={i for i,label in enumerate(labels,1) if re.sub(r'\s+','',label) in ('전체거래금액','현금거래금액','잔액')}
    codecol=ordered.index(22)+1 if 22 in ordered else None
    black=Font(name='맑은 고딕',size=10,color='000000')
    headerfont=Font(name='맑은 고딕',size=10,color='000000',bold=False)
    fill_head=PatternFill('solid',fgColor='E7E9ED')
    styles={}
    for kind in ('입금','지급'):
        cfg=colors.get(kind,{})
        styles[kind]=(
            PatternFill('solid',fgColor=str(cfg.get('bg',DEFAULT_BG[kind])).lstrip('#')),
            Font(name='맑은 고딕',size=10,bold=False,color=str(cfg.get('fg',DEFAULT_FG[kind])).lstrip('#')))
    for row in out.iter_rows(min_row=1,max_row=count+1,max_col=width):
        rn=row[0].row
        is_header=rn==1
        kind=''
        if not is_header and codecol:
            code=out.cell(rn,codecol).value
            if str(code or '').strip().startswith('입금') and colors.get('입금',{}).get('enabled',False):kind='입금'
            elif str(code or '').strip().startswith('지급') and colors.get('지급',{}).get('enabled',True):kind='지급'
        for i,c in enumerate(row,1):
            c.font=(headerfont if is_header else styles[kind][1] if kind else black)
            c.fill=(fill_head if is_header else styles[kind][0] if kind else PatternFill(fill_type=None))
            c.border=Border(left=Side(style='thin' if i==1 else 'dotted',color='989898'),
                            right=Side(style='thin' if i==width else 'dotted',color='989898'),
                            top=Side(style='thin',color='989898'),bottom=Side(style='thin',color='989898'))
            c.alignment=Alignment(
                horizontal=('center' if is_header else 'right' if normalized_header(labels[i-1]) == '전체거래금액' else 'center' if normalized_header(labels[i-1]) in CENTER_ALIGN_HEADERS else 'general'),
                vertical='center', wrap_text=is_header)
            if not is_header and i in money:
                c.number_format='#,##0;-#,##0;0'
        if not is_header:
            out.row_dimensions[rn].height=30
    # 내용 기반 열 너비. 긴 번호열은 출력 폭 제한.
    for i,n in enumerate(ordered,1):
        label=labels[i-1]
        if n in (11,12):max_width=14
        elif n in (26,27):max_width=20
        elif n in (91,92):max_width=13
        elif i in money:max_width=18
        else:max_width=26
        sample_width=max((len(str(out.cell(r,i).value or '').split('\n')[0]) for r in range(2,min(count+2,202))),default=0)
        header_width=max(len(x) for x in str(out.cell(1,i).value).split('\n'))
        computed=max(10,min(max_width,max(header_width*1.8+2,min(sample_width*1.2+2,max_width))))
        out.column_dimensions[get_column_letter(i)].width=computed
    # 비연속 원본열도 출력 좌표 기준으로 그룹 구성.
    sets=[{11,12},{23},{25,26,27,28},{31}]
    for group in sets:
        positions=sorted(i for i,n in enumerate(ordered,1) if n in group)
        if not positions:continue
        chunks=[];chunk=[positions[0]]
        for v in positions[1:]:
            if v==chunk[-1]+1:chunk.append(v)
            else:chunks.append(chunk);chunk=[v]
        chunks.append(chunk)
        for chunk in chunks:
            for idx in chunk:
                dim=out.column_dimensions[get_column_letter(idx)]
                dim.outlineLevel=1
                dim.hidden=True
            if chunk[-1]<width:
                out.column_dimensions[get_column_letter(chunk[-1]+1)].collapsed=True
    out.sheet_properties.outlinePr.summaryRight=True
    out.sheet_view.showOutlineSymbols=True
    out.auto_filter.ref=f'A1:{get_column_letter(width)}{max(1,count+1)}'
    out.sheet_properties.pageSetUpPr.fitToPage=True
    out.page_setup.paperSize=out.PAPERSIZE_A4
    out.page_setup.orientation=out.ORIENTATION_LANDSCAPE
    out.page_setup.fitToWidth=1
    out.page_setup.fitToHeight=0
    out.print_options.horizontalCentered=True
    cm=lambda x:x/2.54
    out.page_margins.left=out.page_margins.right=cm(1)
    out.page_margins.top=out.page_margins.bottom=cm(1)
    out.page_margins.header=out.page_margins.footer=cm(.6)
    out.oddHeader.left.text='&F'
    out.oddFooter.right.text='&N페이지 중 &P페이지'
    out.print_title_rows='1:1'
    # 저장 후 XML 재작성 없이 openpyxl로만 직렬화
    wb.active=0
    # 열려 있는 Excel에서 변환 시트만 활성/선택 상태로 표시
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = (sheet is out)
    if wb.views:
        wb.views[0].activeTab = 0
        wb.views[0].firstSheet = 0
    wb.save(dest)
    # 생성된 파일은 한번 더 열어 구조가 유효한지 확인
    test=load_workbook(dest,read_only=False)
    try:
        if test.sheetnames[:3]!=['변환(요약)','변환(원본칼럼 유지)','원본시트'] or test['변환(요약)'].max_row!=count+1 or test['변환(원본칼럼 유지)'].max_row!=full_count+3:
            raise ValueError('저장된 파일 검증에 실패했습니다.')
    finally:
        test.close()
    return dest,count

class App(WindowBase):
    def __init__(self):
        super().__init__()
        self.title(APP+'  v4.4')
        self.geometry('1020x720'); self.minsize(800,650)
        self.settings=load_settings(); self.items=[]; self.q=queue.Queue(); self.working=False
        self.column_labels={n:colname(n) for n in LOGICAL_COLS}
        self.preset=tk.StringVar(value=self.settings.get('selected',list(self.settings['presets'])[0]))
        self.status=tk.StringVar(value='파일을 추가하세요. 원본 폴더에 _변환.xlsx로 저장하며, 이미 있으면 _변환2.xlsx, _변환3.xlsx로 자동 저장합니다.')
        self.setup_ui(); self.after(120,self.poll)
    def setup_ui(self):
        top=ttk.Frame(self,padding=14);top.pack(fill='both',expand=True)
        ttk.Label(top,text='거래내역 1·2 합치기',font=('맑은 고딕',16)).pack(anchor='w')
        ttk.Label(top,text='변환(요약): 버튼별 열 선택 / 변환(원본칼럼 유지): 원본 열 구조 / 원본시트: 원자료 보존',foreground='#475569').pack(anchor='w',pady=(3,10))
        bar=ttk.Frame(top);bar.pack(fill='x')
        ttk.Button(bar,text='엑셀 파일 추가',command=self.add_files).pack(side='left')
        ttk.Button(bar,text='선택 삭제',command=self.remove_files).pack(side='left',padx=6)
        ttk.Button(bar,text='전체 삭제',command=self.clear_files).pack(side='left')
        ttk.Button(bar,text='환경설정 · 컬럼/색상',command=self.preferences).pack(side='right')
        # 파일 목록을 기존의 약 절반 높이로 줄여 변환 버튼 공간 확보
        box=ttk.Frame(top,height=215);box.pack(fill='both',expand=False,pady=(8,10))
        box.pack_propagate(False)
        self.lst=tk.Listbox(box,selectmode='extended',height=8,font=('맑은 고딕',10))
        scroll=ttk.Scrollbar(box,orient='vertical',command=self.lst.yview)
        self.lst.configure(yscrollcommand=scroll.set)
        self.lst.pack(side='left',fill='both',expand=True);scroll.pack(side='right',fill='y')
        # 빈 파일 목록 안에 드래그 안내문을 겹쳐 표시
        self.drop_hint=tk.Label(self.lst,text='여기로 엑셀 파일 드래그\n또는 위의 「엑셀 파일 추가」 클릭',
                                font=('맑은 고딕',14),fg='#64748B',bg='#F3F6FA',justify='center')
        self.drop_hint.place(relx=.5,rely=.5,anchor='center')
        if DRAG_DROP_AVAILABLE:
            for widget in (self.lst,self.drop_hint):
                widget.drop_target_register(DND_FILES)
                widget.dnd_bind('<<Drop>>',self.on_drop)
        else:
            self.drop_hint.configure(text='엑셀 파일 추가 버튼을 이용하세요\n(드래그 사용 시 tkinterdnd2 설치 필요)')
        ttk.Label(top,text='변환 방식 선택 · 버튼 하나로 목록의 모든 파일 처리',font=('맑은 고딕',10)).pack(anchor='w',pady=(0,5))
        self.button_frame=ttk.Frame(top);self.button_frame.pack(fill='x',pady=(0,8))
        self.refresh_transform_buttons()
        self.progress=ttk.Progressbar(top,mode='determinate');self.progress.pack(fill='x',pady=(7,5))
        ttk.Label(top,textvariable=self.status,wraplength=950).pack(anchor='w')
        # 안내 영역: 화면의 하단 여유 공간에 고정하여 작업 버튼과 분리한다.
        info = ttk.LabelFrame(top, text='프로그램 안내', padding=(12, 9))
        info.pack(side='bottom', fill='x', pady=(16, 0))
        guidance = (
            '1. 입출금내역제공내역 파일을 보기좋게 변환하는 프로그램입니다.\n'
            '2. 업무상 필요한 칼럼들만 선택하여 프리셋 저장 후 활용가능합니다.\n'
            '3. 변환하더라도 원본파일은 유지되므로 안심하고 사용하세요.'
        )
        ttk.Label(info, text=guidance, justify='left', font=('맑은 고딕', 10)).pack(anchor='w')
        ttk.Label(info, text='Made by deepest@shinhan.com', foreground='#475569',
                  font=('맑은 고딕', 9)).pack(anchor='e', pady=(8, 0))
        self.refresh_file_hint()

    def refresh_transform_buttons(self):
        for child in self.button_frame.winfo_children():child.destroy()
        self.transform_buttons=[]
        actions=[('변환(전체)',list(LOGICAL_COLS))]
        for name,cols in self.settings['presets'].items():
            label='변환('+('기본 프리셋1' if name.startswith('기본 프리셋1') else name)+')'
            actions.append((label,list(cols)))
        # 사용자 프리셋이 추가되면 행이 늘어나는 가변 버튼 목록
        for i,(label,cols) in enumerate(actions):
            btn=ttk.Button(self.button_frame,text=label,command=lambda s=cols,n=label:self.start(s,n))
            btn.grid(row=i//4,column=i%4,sticky='ew',padx=3,pady=4,ipady=8)
            self.transform_buttons.append(btn)
        for col in range(4):self.button_frame.columnconfigure(col,weight=1,uniform='transforms')

    def refresh_file_hint(self):
        if self.items:self.drop_hint.place_forget()
        else:self.drop_hint.place(relx=.5,rely=.5,anchor='center')

    def insert_files(self,paths):
        added=[]
        for raw in paths:
            p=str(Path(raw).expanduser())
            if Path(p).suffix.lower()!='.xlsx' or Path(p).name.startswith('~$') or not Path(p).is_file():continue
            if p not in self.items:
                self.items.append(p);self.lst.insert('end',p);added.append(p)
        self.refresh_file_hint()
        if added:
            try:
                labels,_=get_xml_sheet_header(added[0]);self.column_labels=labels
            except Exception as ex:self.status.set('머리글 미리보기 실패: '+str(ex))
        self.status.set(f'{len(self.items)}개 파일 선택됨 (새로 추가 {len(added)}개)')

    def on_drop(self,event):
        # splitlist는 공백 및 중괄호가 있는 Windows 파일 경로를 처리
        try:paths=self.tk.splitlist(event.data)
        except tk.TclError:paths=[event.data]
        self.insert_files(paths)

    def add_files(self):
        paths=filedialog.askopenfilenames(title='복수의 xlsx 파일 선택',filetypes=[('Excel 통합문서','*.xlsx')])
        self.insert_files(paths)

    def remove_files(self):
        for i in reversed(self.lst.curselection()):
            del self.items[i];self.lst.delete(i)
        self.refresh_file_hint()

    def clear_files(self):
        self.items.clear();self.lst.delete(0,'end');self.refresh_file_hint()

    def preferences(self):
        if self.working: return
        win=tk.Toplevel(self);win.title('환경설정 · 표시할 열 및 색상');win.geometry(f'1160x{min(835,max(660,win.winfo_screenheight()-115))}');win.minsize(1040,630);win.transient(self);win.grab_set()
        body=ttk.Frame(win,padding=9);body.pack(fill='both',expand=True)
        note='요약 시트의 표시 열만 선택합니다. 전체 시트는 원본 A~CL 90개 열을 그대로 유지합니다.'
        ttk.Label(body,text=note,wraplength=1080).pack(anchor='w',pady=(0,5))
        row=ttk.Frame(body);row.pack(fill='x')
        ttk.Label(row,text='프리셋:').pack(side='left')
        curr=tk.StringVar(value=self.preset.get())
        combo=ttk.Combobox(row,textvariable=curr,values=list(self.settings['presets']),state='readonly',width=38)
        combo.pack(side='left',padx=8)
        default_preset='기본 프리셋1 (거래내역만-대출제외)'
        def delete_preset():
            name=curr.get()
            if name==default_preset:
                messagebox.showwarning('기본 프리셋 보호','기본 프리셋은 삭제할 수 없습니다.',parent=win)
                return
            if name not in self.settings['presets']:
                messagebox.showwarning('프리셋 확인','삭제할 프리셋을 선택하세요.',parent=win)
                return
            if not messagebox.askyesno('프리셋 삭제',f'“{name}” 프리셋을 삭제하시겠습니까?\n이 작업은 되돌릴 수 없습니다.',parent=win):
                return
            del self.settings['presets'][name]
            curr.set(default_preset)
            self.preset.set(default_preset)
            combo['values']=list(self.settings['presets'])
            apply_list(self.settings['presets'][default_preset])
            self.persist()
            messagebox.showinfo('삭제 완료',f'“{name}” 프리셋을 삭제했습니다.',parent=win)
        ttk.Button(row,text='선택 프리셋 삭제',command=delete_preset).pack(side='left',padx=(2,8))
        frame=ttk.LabelFrame(body,text='변환(요약)에 표시할 열 (체크 = 포함)',padding=3)
        frame.pack(fill='both',expand=True,pady=(6,5))
        # 58개 원본 열을 읽는 순서대로 4개 패널에 균등 배치.
        # 원본의 열 번호를 그대로 표기하여 선택/제외 프리셋과 호환.
        vars={}
        selected=set(self.settings['presets'].get(curr.get(),DEFAULT_COLS))
        sections=[LOGICAL_COLS[i:i+15] for i in range(0,len(LOGICAL_COLS),15)]
        for panel_index,section in enumerate(sections):
            panel=ttk.Frame(frame,padding=(3,1))
            panel.grid(row=0,column=panel_index,sticky='nsew',padx=2)
            title=f'{colname(section[0])} ~ {colname(section[-1])}'
            ttk.Label(panel,text=title,font=('맑은 고딕',10)).pack(anchor='w',pady=(0,3))
            for n in section:
                v=tk.BooleanVar(value=n in selected); vars[n]=v
                desc=str(LOAN_LABELS.get(n,self.column_labels.get(n,colname(n)))).replace('\n',' ').strip()
                note2='' if 43<=n<=53 else (' [담당자]' if 86<=n<=90 else '')
                ttk.Checkbutton(panel,text=f'{colname(n):>2}  {desc}{note2}',variable=v).pack(anchor='w',pady=0)
        for panel_index in range(4):
            frame.columnconfigure(panel_index,weight=1,uniform='column')
        def apply_list(cols):
            for n,v in vars.items():v.set(n in cols)
        def current_cols(): return [n for n in LOGICAL_COLS if vars[n].get()]
        def from_preset(evt=None): apply_list(self.settings['presets'].get(curr.get(),DEFAULT_COLS))
        combo.bind('<<ComboboxSelected>>',from_preset)
        tools=ttk.Frame(body);tools.pack(fill='x')
        ttk.Button(tools,text='전체 선택',command=lambda:apply_list(LOGICAL_COLS)).pack(side='left',padx=2)
        ttk.Button(tools,text='전체 해제',command=lambda:apply_list([])).pack(side='left',padx=2)
        ttk.Button(tools,text='대출 AQ~BA 제외',command=lambda:apply_list([n for n in current_cols() if not 43<=n<=53])).pack(side='left',padx=2)
        ttk.Button(tools,text='법적근거·담당자 CH~CL 제외',command=lambda:apply_list([n for n in current_cols() if not 86<=n<=90])).pack(side='left',padx=2)
        ttk.Button(tools,text='기본 K,L,V~AM',command=lambda:apply_list(DEFAULT_COLS)).pack(side='left',padx=2)
        colsframe=ttk.LabelFrame(body,text='입금/지급 거래 행 전체 색상',padding=4)
        colsframe.pack(fill='x',pady=(4,5))
        tempcolors=copy.deepcopy(self.settings['colors'])
        for kind in ('입금','지급'):
            tempcolors.setdefault(kind, {})
            tempcolors[kind].setdefault('enabled',kind=='지급')
            tempcolors[kind].setdefault('bg',DEFAULT_BG[kind])
            tempcolors[kind].setdefault('fg',DEFAULT_FG[kind])
        color_choice=tk.StringVar(value=self.settings.get('color_preset',DEFAULT_COLOR_PRESET))
        if color_choice.get() not in COLOR_PRESETS: color_choice.set('사용자 지정')
        preset_row=ttk.Frame(colsframe);preset_row.pack(fill='x',pady=(0,4))
        ttk.Label(preset_row,text='색상 프리셋').pack(side='left')
        color_combo=ttk.Combobox(preset_row,textvariable=color_choice,
            values=list(COLOR_PRESETS)+['사용자 지정'],state='readonly',width=35)
        color_combo.pack(side='left',padx=10)
        ttk.Label(preset_row,text='선택 즉시 미리보기 반영',foreground='#64748B').pack(side='left')
        preview_info={}
        color_buttons={}
        def refresh_all():
            for kind in ('입금','지급'):
                pv,flag=preview_info[kind]
                enabled=bool(tempcolors[kind].get('enabled',False))
                flag.set(enabled)
                pv.configure(bg=tempcolors[kind]['bg'] if enabled else '#FFFFFF',
                             fg=tempcolors[kind]['fg'] if enabled else '#000000',
                             text=f'  {kind} 거래 행 샘플  ' + ('' if enabled else '(색상 없음)'))
                bgbtn,fgbtn=color_buttons[kind]
                bgbtn.configure(bg=tempcolors[kind]['bg'], fg='#111111', activebackground=tempcolors[kind]['bg'],
                                text='배경색', relief='raised')
                fgbtn.configure(bg='#FFFFFF', fg=tempcolors[kind]['fg'], activeforeground=tempcolors[kind]['fg'],
                                text='글자색', relief='raised')
        def use_color_preset(evt=None):
            name=color_choice.get()
            if name in COLOR_PRESETS:
                tempcolors.clear();tempcolors.update(copy.deepcopy(COLOR_PRESETS[name]));refresh_all()
        color_combo.bind('<<ComboboxSelected>>',use_color_preset)
        for kind in ('입금','지급'):
            rr=ttk.Frame(colsframe);rr.pack(fill='x',pady=2)
            enabled=tk.BooleanVar(value=tempcolors[kind]['enabled'])
            def on_toggle(kind=kind,enabled=enabled):
                tempcolors[kind]['enabled']=enabled.get()
                color_choice.set('사용자 지정');refresh_all()
            ttk.Checkbutton(rr,text=f'{kind} 색상 적용',variable=enabled,command=on_toggle).pack(side='left',padx=(0,8))
            preview=tk.Label(rr,text='',font=('맑은 고딕',10),relief='groove',width=24)
            preview.pack(side='right',padx=8)
            preview_info[kind]=(preview,enabled)
            btns={}
            for which,textlabel in (('bg','배경색'),('fg','글자색')):
                def choose(kind=kind,which=which):
                    color=colorchooser.askcolor(color=tempcolors[kind][which],parent=win,title=f'{kind} {textlabel}')[1]
                    if color:
                        tempcolors[kind][which]=color
                        color_choice.set('사용자 지정')
                        refresh_all()
                btn=tk.Button(rr,text=textlabel,command=choose, width=9, font=('맑은 고딕',9),
                              relief='raised', borderwidth=2, cursor='hand2')
                btn.pack(side='left',padx=3)
                btns[which]=btn
            color_buttons[kind]=(btns['bg'],btns['fg'])
        refresh_all()
        buttons=ttk.Frame(body);buttons.pack(fill='x')
        def do_save():
            cols=current_cols()
            if not cols: messagebox.showwarning('알림','하나 이상의 열을 선택하세요.',parent=win);return
            import tkinter.simpledialog as sd
            name=sd.askstring('프리셋 저장','저장할 프리셋 이름:',initialvalue='사용자 프리셋',parent=win)
            if not name: return
            if name=='기본 프리셋1 (거래내역만-대출제외)':
                messagebox.showwarning('기본값 보호','기본 프리셋은 수정할 수 없습니다. 다른 이름을 사용하세요.',parent=win);return
            self.settings['presets'][name]=cols;curr.set(name)
            combo['values']=list(self.settings['presets'])
            self.settings['colors']=copy.deepcopy(tempcolors);self.settings['color_preset']=color_choice.get();self.preset.set(name)
            self.persist();messagebox.showinfo('저장 완료',f'“{name}” 프리셋 저장 완료',parent=win)
        def commit():
            cols=current_cols()
            if not cols: messagebox.showwarning('알림','하나 이상의 열을 선택하세요.',parent=win);return
            name=curr.get()
            if cols != self.settings['presets'].get(name):
                name='임시 선택'
                self.settings['presets'][name]=cols
            self.settings['colors']=copy.deepcopy(tempcolors);self.settings['color_preset']=color_choice.get();self.preset.set(name);self.persist();win.destroy()
        ttk.Button(buttons,text='현재 선택을 새 프리셋으로 저장',command=do_save).pack(side='left')
        ttk.Button(buttons,text='적용 후 닫기',command=commit).pack(side='right')
        ttk.Button(buttons,text='취소',command=win.destroy).pack(side='right',padx=8)
    def persist(self):
        self.settings['selected']=self.preset.get()
        save_settings(self.settings)
        self.refresh_transform_buttons()
    def start(self,selected,action_name):
        if self.working:return
        if not self.items: messagebox.showwarning('파일 필요','변환할 파일을 하나 이상 추가하세요.');return
        if not selected: messagebox.showwarning('열 필요','표시할 열이 없습니다.');return
        self.persist();self.working=True
        for btn in self.transform_buttons:btn.configure(state='disabled')
        self.status.set(f'{action_name} 처리 중...')
        self.progress['maximum']=len(self.items);self.progress['value']=0
        paths=list(self.items); colors=json.loads(json.dumps(self.settings['colors']));skip=True
        def work():
            ok=[];errors=[]
            for i,p in enumerate(paths):
                try:
                    dest,count=process_file(p,selected,colors,skip,lambda s:self.q.put(('status',s)))
                    ok.append(f'{dest.name} ({count:,}건)')
                except Exception as exc:
                    errors.append(f'{Path(p).name}: {exc}')
                self.q.put(('progress',i+1))
            self.q.put(('done',(ok,errors)))
        threading.Thread(target=work,daemon=True).start()
    def poll(self):
        try:
            while True:
                key,val=self.q.get_nowait()
                if key=='status':self.status.set(val)
                elif key=='progress':self.progress['value']=val
                elif key=='done':
                    ok,errors=val;self.working=False
                    for btn in self.transform_buttons:btn.configure(state='normal')
                    self.status.set(f'완료 {len(ok)}개 / 실패 {len(errors)}개')
                    msg=f'변환 성공 {len(ok)}개 / 실패 {len(errors)}개\n\n'+ '\n'.join(ok[:12])
                    if errors:msg+='\n\n실패 내역:\n'+'\n'.join(errors[:12])
                    messagebox.showwarning('일괄 변환 결과',msg) if errors else messagebox.showinfo('일괄 변환 완료',msg)
        except queue.Empty:pass
        self.after(120,self.poll)

if __name__=='__main__':
    App().mainloop()
